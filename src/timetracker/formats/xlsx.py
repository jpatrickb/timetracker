# formats/xlsx.py
#
# Author: Patrick Beal
#
# Invoice as an Excel workbook, following docs/invoice-example/

from datetime import timedelta
from pathlib import Path

import pendulum

from timetracker.errors import TimeTrackerError
from timetracker.formats import header, range_label, total_row
from timetracker.invoice import InvoiceDocument, period_label
from timetracker.report import Report
from timetracker.user import address_lines, full_name

HEADERS = ("Date", "Project", "Hours", "Amount", "Description")

DATE_FORMAT = "M/d/yyyy"
HOURS_FORMAT = "[h]:mm:ss"  # counts past 24 hours instead of wrapping
MONEY_FORMAT = '"$"#,##0.00'

# The palette and metrics of docs/invoice-example/
INK = "223642"  # the dark navy of the banner and the table's rules
PAPER = "FFFFFF"
STRIPE = "F6F8F9"  # every second table row
BANNER_FONT = "Roboto"
TABLE_FONT = "Arial"
TABLE_ROW_HEIGHT = 22.5
TITLE_LINE_HEIGHT = 24  # the banner's 18pt text, with its leading
COLUMNS = "BCDEF"  # A and G are the left and right margins
WIDTHS = {"D": 13.25, "F": 46.25}  # the rest keep the default width
DEFAULT_WIDTH = 12.63


def render_xlsx(document: InvoiceDocument, path: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.properties import PageSetupProperties
    except ImportError:
        raise TimeTrackerError(
            "Excel output needs openpyxl. Reinstall the app with `uv sync`."
        ) from None

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Invoice"

    banner_fill = PatternFill("solid", fgColor=INK)
    stripe_fill = PatternFill("solid", fgColor=STRIPE)
    paper_fill = PatternFill("solid", fgColor=PAPER)
    rule = Side(style="thick", color=INK)
    thin = Side(style="thin", color=INK)
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    white = Font(name=BANNER_FONT, color=PAPER)
    label = Font(name=BANNER_FONT, bold=True, size=12)

    # --- Banner ---
    title = f"HOURLY INVOICE\n{period_label(document)}"
    if document.predicted:
        title += "\n(DRAFT)"
    sheet.merge_cells("B1:D1")
    sheet["B1"] = title
    sheet["B1"].font = Font(name=BANNER_FONT, bold=True, size=18, color=PAPER)
    sheet["B1"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["E1"] = "INVOICE NUMBER:"
    sheet["E1"].alignment = Alignment(horizontal="right", vertical="top")
    sheet["F1"] = document.number
    sheet["F1"].alignment = Alignment(horizontal="left", vertical="top")
    for column in "ABCDEFG":
        cell = sheet[f"{column}1"]
        cell.fill = banner_fill
        cell.border = Border(bottom=rule)
        if cell.font.color is None or cell.font.color.rgb != f"00{PAPER}":
            cell.font = white
    # Excel won't autofit a merged cell, so the banner is sized to its lines
    sheet.row_dimensions[1].height = TITLE_LINE_HEIGHT * title.count("\n") + 30

    # --- Address blocks, each headed by a labeled rule ---
    row_number = 3
    row_number = _block(sheet, row_number, "From:", _sender(document), label, rule)
    row_number = _block(
        sheet, row_number, "Bill To:", [document.client_name], label, rule
    )

    # --- Table ---
    header_row = row_number
    for column, heading in zip(COLUMNS, HEADERS):
        cell = sheet[f"{column}{header_row}"]
        cell.value = heading
        cell.font = Font(name=TABLE_FONT, bold=True, color=PAPER)
        cell.fill = banner_fill
        cell.border = grid
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[header_row].height = TABLE_ROW_HEIGHT

    for offset, line in enumerate(document.rows, start=1):
        row_number = header_row + offset
        sheet[f"B{row_number}"] = line.day
        sheet[f"C{row_number}"] = line.project or ""
        sheet[f"D{row_number}"] = timedelta(seconds=line.seconds)
        sheet[f"E{row_number}"] = _amount(line, row_number)
        sheet[f"F{row_number}"] = line.description or ""
        # Stripes start on the paper shade, as in the example
        shade = paper_fill if offset % 2 else stripe_fill
        _style_table_row(sheet, row_number, grid, shade)

    total_row_number = header_row + len(document.rows) + 1
    first = header_row + 1
    last = total_row_number - 1
    sheet[f"B{total_row_number}"] = "Total"
    sheet[f"D{total_row_number}"] = f"=SUM(D{first}:D{last})"
    sheet[f"E{total_row_number}"] = f"=SUM(E{first}:E{last})"
    _style_table_row(sheet, total_row_number, grid, paper_fill)
    for column in COLUMNS:
        cell = sheet[f"{column}{total_row_number}"]
        cell.font = Font(name=TABLE_FONT, bold=True)

    if document.user["payment_notes"]:
        notes_row = total_row_number + 2
        sheet[f"B{notes_row}"] = document.user["payment_notes"]
        sheet[f"B{notes_row}"].font = Font(name=BANNER_FONT)

    # --- Sheet setup ---
    sheet.sheet_view.showGridLines = False
    sheet.sheet_format.defaultColWidth = DEFAULT_WIDTH
    for column, width in WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _sender(document: InvoiceDocument) -> list[str]:
    return [full_name(document.user), *address_lines(document.user)]


def _block(sheet, row_number: int, heading: str, lines: list[str], label, rule) -> int:
    """
    One address block: a bold label, its lines, and a rule under each, like
    the example's. Returns the first free row two below it.
    """
    from openpyxl.styles import Font, PatternFill

    sheet[f"B{row_number}"] = heading
    sheet[f"B{row_number}"].font = label
    sheet[f"B{row_number}"].fill = PatternFill("solid", fgColor=PAPER)
    _rule_across(sheet, row_number, rule)

    for offset, line in enumerate(lines, start=1):
        cell = sheet[f"B{row_number + offset}"]
        cell.value = line
        # The name leads the block, so it carries the weight
        cell.font = Font(name=BANNER_FONT, bold=offset == 1)

    last = row_number + len(lines)
    _rule_across(sheet, last, rule)
    return last + 2


def _rule_across(sheet, row_number: int, rule):
    from openpyxl.styles import Border

    for column in COLUMNS:
        sheet[f"{column}{row_number}"].border = Border(bottom=rule)


def _style_table_row(sheet, row_number: int, grid, fill):
    from openpyxl.styles import Alignment, Font

    sheet.row_dimensions[row_number].height = TABLE_ROW_HEIGHT
    formats = {"B": DATE_FORMAT, "D": HOURS_FORMAT, "E": MONEY_FORMAT}
    for column in COLUMNS:
        cell = sheet[f"{column}{row_number}"]
        cell.font = Font(name=TABLE_FONT)
        cell.fill = fill
        cell.border = grid
        cell.alignment = Alignment(vertical="center", wrap_text=column == "F")
        if column in formats:
            cell.number_format = formats[column]


def _amount(line, row_number: int):
    """
    The example shows its arithmetic, so a known rate becomes a formula over
    the hours cell. Rows billed at mixed rates fall back to their total.
    """
    if line.rate is not None:
        return f"=D{row_number}*24*{line.rate:g}"
    return None if line.cents is None else line.cents / 100


# --- Reports ---

REPORT_FORMATS = {
    "start": "M/d/yyyy h:mm:ss",
    "end": "M/d/yyyy h:mm:ss",
    "duration": HOURS_FORMAT,
    "rate": MONEY_FORMAT,
    "pay": MONEY_FORMAT,
}


def render_report_xlsx(report: Report, path: Path):
    """A report as a spreadsheet, with real dates, durations, and money."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise TimeTrackerError(
            "Excel output needs openpyxl. Reinstall the app with `uv sync`."
        ) from None

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Report"
    bold = Font(bold=True)

    sheet["A1"] = "TIME REPORT"
    sheet["A1"].font = bold
    sheet["A2"] = range_label(report)

    header_row = 4
    for column, field in enumerate(report.fields, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header(report, field))
        cell.font = bold

    for offset, row in enumerate(report.rows, start=1):
        _write_report_row(sheet, header_row + offset, report, row)

    total_line = header_row + len(report.rows) + 1
    _write_report_row(sheet, total_line, report, total_row(report))
    for column in range(1, len(report.fields) + 1):
        sheet.cell(row=total_line, column=column).font = bold

    for column, field in enumerate(report.fields, start=1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = 60 if field == "description" else 20

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_report_row(sheet, row_number: int, report: Report, row: dict):
    from openpyxl.styles import Alignment

    for column, field in enumerate(report.fields, start=1):
        cell = sheet.cell(row=row_number, column=column)
        cell.value = _xlsx_value(field, row[field])
        if row[field] is not None and field in REPORT_FORMATS:
            cell.number_format = REPORT_FORMATS[field]
        if field == "description":
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _xlsx_value(field: str, value):
    """Numbers, dates, and durations stay typed so Excel can work with them."""
    if value is None:
        return None
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    if field in ("start", "end") and isinstance(value, int):
        return pendulum.from_timestamp(value, tz="local").naive()
    if field == "duration" and isinstance(value, int):
        return timedelta(seconds=value)
    if field == "pay" and isinstance(value, int):
        return value / 100
    return value
