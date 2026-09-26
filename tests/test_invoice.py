from datetime import date

import openpyxl
import pendulum
import pytest
import time_machine
from typer.testing import CliRunner

from timetracker import clients, clock, db, invoice, user
from timetracker.cli import app
from timetracker.errors import TimeTrackerError
from timetracker.formats.pdf import _typst_source, render_pdf
from timetracker.formats.xlsx import render_xlsx

NOW = pendulum.local(2026, 9, 29, 12, 0, 0)

DETAILS = {
    "first_name": "Patrick",
    "last_name": "Beal",
    "address_line_1": "802 N 700 E #4",
    "city": "Provo",
    "state": "UT",
    "zip_code": "84606",
}


@pytest.fixture(autouse=True)
def frozen_time():
    with time_machine.travel(NOW, tick=False):
        yield


@pytest.fixture
def conn(tmp_path):
    with db.connection(tmp_path / "test.db") as conn:
        user.save_user(conn, **DETAILS)
        clients.create_client(conn, "TechForce Advisors", pay_rate=35, aliases=["TFA"])
        clients.create_client(conn, "Patea", pay_rate=60)
        clients.create_project(conn, 1, "Extraction")
        clients.create_project(conn, 1, "Pro Forma", pay_rate=45)
        yield conn


def add(conn, start, end, **kwargs):
    entry_id, _ = clock.add_entry(
        conn, clock.parse_time(start), clock.parse_time(end), **kwargs
    )
    return entry_id


def draft(conn, entry_ids, client_id=1):
    return invoice.create_draft(
        conn, client_id, date(2026, 9, 1), date(2026, 9, 14), entry_ids
    )


# --- User details ---


def test_save_user_merges_fields(conn):
    user.save_user(conn, city="Orem", email="p@example.com")
    saved = user.require_user(conn)
    assert (saved["city"], saved["state"]) == ("Orem", "UT")
    assert saved["email"] == "p@example.com"


def test_save_user_requires_the_basics(tmp_path):
    with db.connection(tmp_path / "empty.db") as conn:
        with pytest.raises(TimeTrackerError, match="Still missing: ZIP code"):
            user.save_user(conn, **{**DETAILS, "zip_code": ""})
        with pytest.raises(TimeTrackerError, match="Run `tt setup`"):
            user.require_user(conn)


def test_address_lines_skip_blanks(conn):
    assert user.address_lines(user.require_user(conn)) == [
        "802 N 700 E #4",
        "Provo, UT 84606",
    ]
    user.save_user(conn, address_line_2="Apt 2", phone="555-0100")
    assert user.address_lines(user.require_user(conn)) == [
        "802 N 700 E #4",
        "Apt 2",
        "Provo, UT 84606",
        "555-0100",
    ]


# --- Drafts ---


def test_draft_holds_no_number_but_predicts_one(conn):
    entry = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    draft(conn, [entry])
    listed = invoice.list_invoices(conn)[0]
    assert listed["status"] == "draft"
    assert (listed["predicted"], listed["invoice_number"]) == (True, 1)
    assert conn.execute("SELECT invoice_number FROM invoices").fetchone()[0] is None


def test_entry_can_sit_on_two_drafts(conn):
    entry = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    draft(conn, [entry])
    draft(conn, [entry])
    assert len(invoice.list_invoices(conn)) == 2


def test_draft_rejects_open_and_foreign_entries(conn):
    open_id, _ = clock.clock_in(conn, clock.parse_time("09:00"), client="TFA")
    with pytest.raises(TimeTrackerError, match="open entries"):
        draft(conn, [open_id])

    other = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="Patea")
    with pytest.raises(TimeTrackerError, match="different client"):
        draft(conn, [other])


def test_draft_needs_entries(conn):
    with pytest.raises(TimeTrackerError, match="No entries"):
        draft(conn, [])


# --- Issuing, voiding, deleting ---


def test_issue_allocates_number_per_client(conn):
    tfa = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    patea = add(conn, "2026-09-01 11:00", "2026-09-01 12:00", client="Patea")
    assert invoice.issue(conn, draft(conn, [tfa])) == 1
    # Each client has its own sequence, so Patea also starts at 1
    assert invoice.issue(conn, draft(conn, [patea], client_id=2)) == 1
    tfa2 = add(conn, "2026-09-02 09:00", "2026-09-02 10:00", client="TFA")
    assert invoice.issue(conn, draft(conn, [tfa2])) == 2


def test_issue_blocks_double_billing_and_names_entries(conn):
    entry = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    first, second = draft(conn, [entry]), draft(conn, [entry])
    invoice.issue(conn, first)
    with pytest.raises(TimeTrackerError, match=r"entry 1 \(on #1\)"):
        invoice.issue(conn, second)


def test_void_frees_entries_and_keeps_number(conn):
    entry = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    first, second = draft(conn, [entry]), draft(conn, [entry])
    invoice.issue(conn, first)
    invoice.void(conn, first)
    # The number stays reserved, so the next issued invoice is #2
    assert invoice.issue(conn, second) == 2


def test_status_changes_are_restricted(conn):
    entry = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    invoice_id = draft(conn, [entry])
    with pytest.raises(TimeTrackerError, match="Only issued invoices can be voided"):
        invoice.void(conn, invoice_id)

    invoice.issue(conn, invoice_id)
    with pytest.raises(TimeTrackerError, match="already issued"):
        invoice.issue(conn, invoice_id)
    with pytest.raises(TimeTrackerError, match="Only drafts can be deleted"):
        invoice.delete_draft(conn, invoice_id)


def test_deleting_a_draft_keeps_the_entries(conn):
    entry = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", client="TFA")
    invoice.delete_draft(conn, draft(conn, [entry]))
    assert invoice.list_invoices(conn) == []
    assert clock.get_entry(conn, entry) is not None


def test_missing_invoice(conn):
    with pytest.raises(TimeTrackerError, match="No invoice with ID 99"):
        invoice.build_document(conn, 99)


# --- Document ---


@pytest.fixture
def document(conn):
    entries = [
        add(
            conn,
            "2026-09-01 09:00",
            "2026-09-01 11:57:47",
            project="Extraction",
            description="Extraction work",
        ),
        add(
            conn,
            "2026-09-02 13:00",
            "2026-09-02 13:52:46",
            project="Extraction",
            description="More extraction",
        ),
        add(
            conn,
            "2026-09-02 22:30",
            "2026-09-03 00:45",
            project="Pro Forma",
            description="Late modeling",
        ),
    ]
    return invoice.build_document(conn, draft(conn, entries))


def test_document_rows_are_one_project_per_day(document):
    assert [(r.day, r.project, r.seconds, r.cents) for r in document.rows] == [
        (date(2026, 9, 1), "Extraction", 10667, 10371),
        (date(2026, 9, 2), "Extraction", 3166, 3078),
        (date(2026, 9, 2), "Pro Forma", 5400, 6750),
        (date(2026, 9, 3), "Pro Forma", 2700, 3375),
    ]
    # The split session's halves add back up to 2.25 hours at $45
    assert document.rows[2].cents + document.rows[3].cents == 10125


def test_document_totals_match_its_rows(document):
    assert document.total_seconds == sum(r.seconds for r in document.rows)
    assert document.total_cents == sum(r.cents for r in document.rows)
    assert document.total_cents == 23574


def test_document_period_and_number(document):
    assert (document.period_start, document.period_end) == (
        date(2026, 9, 1),
        date(2026, 9, 14),
    )
    assert (document.number, document.predicted) == (1, True)
    assert invoice.period_label(document) == "Sep 1 - Sep 14"


def test_filename_follows_the_spec(document):
    assert invoice.filename(document, "pdf") == (
        "techforce-advisors-2026-09-01-2026-09-14-hourly-invoice-patrick-beal-no-1.pdf"
    )


def test_regenerating_uses_linked_entries_not_filters(conn, document):
    # A later entry in the same period doesn't appear on an existing invoice
    add(conn, "2026-09-04 09:00", "2026-09-04 10:00", client="TFA")
    again = invoice.build_document(conn, document.invoice_id)
    assert [r.day for r in again.rows] == [r.day for r in document.rows]


# --- Rendered files ---


def test_xlsx_layout(document, tmp_path):
    path = tmp_path / "invoice.xlsx"
    render_xlsx(document, path)
    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None

    assert sheet["B1"].value == "HOURLY INVOICE\nSep 1 - Sep 14\n(DRAFT)"
    assert sheet["F1"].value == 1
    assert sheet["B4"].value == "Patrick Beal"

    # The table starts below the address block, whose height varies
    header_row = next(
        row for row in range(1, sheet.max_row + 1) if sheet[f"B{row}"].value == "Date"
    )
    assert [sheet[f"{c}{header_row}"].value for c in "BCDEF"] == [
        "Date",
        "Project",
        "Hours",
        "Amount",
        "Description",
    ]
    first = header_row + 1
    # Amounts show their arithmetic off the hours cell, as the example does
    assert sheet[f"E{first}"].value == f"=D{first}*24*35"
    assert sheet[f"D{first}"].number_format == "[h]:mm:ss"
    assert sheet[f"B{first}"].number_format == "M/d/yyyy"
    assert sheet[f"E{first}"].number_format == '"$"#,##0.00'

    total_row = header_row + 1 + len(document.rows)
    last = total_row - 1
    assert sheet[f"B{total_row}"].value == "Total"
    assert sheet[f"D{total_row}"].value == f"=SUM(D{first}:D{last})"
    assert sheet[f"E{total_row}"].value == f"=SUM(E{first}:E{last})"


def test_xlsx_matches_the_example_styling(document, tmp_path):
    path = tmp_path / "invoice.xlsx"
    render_xlsx(document, path)
    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None

    # The banner spans the full width in the example's navy, with white text
    assert sheet.sheet_view.showGridLines is False
    assert sheet["A1"].fill.fgColor.rgb == "00223642"
    assert sheet["G1"].fill.fgColor.rgb == "00223642"
    assert sheet["B1"].font.color.rgb == "00FFFFFF"
    assert sheet["B1"].font.size == 18

    header_row = next(
        row for row in range(1, sheet.max_row + 1) if sheet[f"B{row}"].value == "Date"
    )
    # The table header takes the navy too, and every cell is ruled thin
    assert sheet[f"B{header_row}"].fill.fgColor.rgb == "00223642"
    assert sheet[f"B{header_row}"].border.bottom.style == "thin"
    assert sheet.row_dimensions[header_row].height == 22.5

    # Rows alternate white and the example's off-white stripe
    assert sheet[f"B{header_row + 1}"].fill.fgColor.rgb == "00FFFFFF"
    assert sheet[f"B{header_row + 2}"].fill.fgColor.rgb == "00F6F8F9"


def test_xlsx_falls_back_to_a_value_when_a_row_mixes_rates(conn, tmp_path):
    first = add(conn, "2026-09-01 09:00", "2026-09-01 10:00", project="Extraction")
    second = add(conn, "2026-09-01 11:00", "2026-09-01 12:00", project="Extraction")
    # A re-rated entry lands on the same row as one billed at the old rate
    with conn:
        conn.execute(
            "UPDATE time_entries SET pay_rate_hourly = 50 WHERE entry_id = ?", (second,)
        )
    document = invoice.build_document(conn, draft(conn, [first, second]))

    (line,) = document.rows
    assert line.rate is None  # $35 and $50 can't share one formula

    path = tmp_path / "invoice.xlsx"
    render_xlsx(document, path)
    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None
    header_row = next(
        row for row in range(1, sheet.max_row + 1) if sheet[f"B{row}"].value == "Date"
    )
    assert sheet[f"E{header_row + 1}"].value == 85.0


def test_pdf_is_written(document, tmp_path):
    path = tmp_path / "invoice.pdf"
    render_pdf(document, path)
    assert path.read_bytes().startswith(b"%PDF")


def test_pdf_escapes_typst_markup(conn):
    entry = add(
        conn,
        "2026-09-01 09:00",
        "2026-09-01 10:00",
        client="TFA",
        description="Fixed #4 and *starred* $math$",
    )
    document = invoice.build_document(conn, draft(conn, [entry]))
    source = _typst_source(document)
    assert "\\#4" in source
    assert "\\*starred\\*" in source


# --- CLI ---


@pytest.fixture
def cli(tmp_path, monkeypatch):
    monkeypatch.setenv("TIMETRACKER_DB", str(tmp_path / "data" / "cli.db"))
    monkeypatch.setattr(db, "CONFIG_PATH", tmp_path / "config.toml")
    runner = CliRunner()
    invoke = lambda *args, input=None: runner.invoke(app, list(args), input=input)  # noqa: E731
    return invoke


def setup_cli(cli):
    cli("client", "add", "Acme", "--pay-rate-hourly", "40")
    cli(
        "add",
        "--start-time",
        "2026-09-01 08:00",
        "--end-time",
        "2026-09-01 09:00",
        "--client",
        "Acme",
    )


def test_cli_invoice_needs_user_details_and_leaves_no_draft(cli):
    setup_cli(cli)
    result = cli("report", "--start", "2026-09-01", "--end", "2026-09-14", "--invoice")
    assert result.exit_code == 1
    assert "tt setup" in result.output
    # No draft was left behind by the failed run
    assert "Acme" not in cli("invoice", "list").output


def test_cli_setup_then_invoice_cycle(cli, tmp_path):
    answers = "\n".join(
        [
            str(tmp_path / "data" / "cli.db"),
            "Patrick",
            "Beal",
            "802 N 700 E #4",
            "",
            "Provo",
            "UT",
            "84606",
            "",
            "",
            "",
        ]
    )
    assert cli("setup", input=answers + "\n").exit_code == 0
    setup_cli(cli)

    made = cli("report", "--start", "2026-09-01", "--end", "2026-09-14", "--invoice")
    assert made.exit_code == 0
    assert "Created draft invoice 1" in made.output
    saved = tmp_path / "data" / "invoices"
    assert [p.name for p in saved.iterdir()] == [
        "acme-2026-09-01-2026-09-14-hourly-invoice-patrick-beal-no-1.pdf"
    ]

    issued = cli("invoice", "issue", "1")
    assert "Issued invoice #1" in issued.output

    # Billed entries are excluded, so a second invoice finds nothing
    again = cli("report", "--start", "2026-09-01", "--end", "2026-09-14", "--invoice")
    assert again.exit_code == 1
    assert "No billable entries" in again.output

    forced = cli(
        "report",
        "--start",
        "2026-09-01",
        "--end",
        "2026-09-14",
        "--invoice",
        "--include-billed",
        "--output",
        "xlsx",
    )
    assert forced.exit_code == 0
    assert (
        saved / "acme-2026-09-01-2026-09-14-hourly-invoice-patrick-beal-no-2.xlsx"
    ).exists()


def test_cli_invoice_needs_one_client(cli, tmp_path):
    cli(
        "setup",
        input="\n".join(
            [
                str(tmp_path / "data" / "cli.db"),
                "P",
                "B",
                "1 St",
                "",
                "Provo",
                "UT",
                "84606",
                "",
                "",
                "",
            ]
        )
        + "\n",
    )
    setup_cli(cli)
    cli("client", "add", "Other", "--pay-rate-hourly", "10")
    cli(
        "add",
        "--start-time",
        "2026-09-02 08:00",
        "--end-time",
        "2026-09-02 09:00",
        "--client",
        "Other",
    )

    result = cli("report", "--start", "2026-09-01", "--end", "2026-09-14", "--invoice")
    assert result.exit_code == 1
    assert "one client" in result.output


def test_cli_include_billed_without_invoice_errors(cli):
    setup_cli(cli)
    result = cli("report", "--include-billed")
    assert result.exit_code == 1
    assert "only applies with --invoice" in result.output


def test_cli_user_edit_and_show(cli, tmp_path):
    cli(
        "setup",
        input="\n".join(
            [
                str(tmp_path / "data" / "cli.db"),
                "P",
                "B",
                "1 St",
                "",
                "Provo",
                "UT",
                "84606",
                "",
                "",
                "",
            ]
        )
        + "\n",
    )
    assert cli("user", "edit", "--city", "Orem").exit_code == 0
    assert "Orem" in cli("user", "show").output
    assert cli("user", "edit").exit_code == 1
