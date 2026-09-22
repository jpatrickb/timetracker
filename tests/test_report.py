import csv
import io
import json
from datetime import date, datetime, timedelta

import openpyxl

import pendulum
import pytest
import time_machine
from typer.testing import CliRunner

from timetracker import clients, clock, db, report
from timetracker.cli import app
from timetracker.errors import TimeTrackerError
from timetracker.formats.delimited import render_csv, render_tsv
from timetracker.formats.json_format import render_json
from timetracker.formats.table import render_markdown

# A Tuesday. This week runs Mon Sep 28 to Sun Oct 4.
NOW = pendulum.local(2026, 9, 29, 12, 0, 0)


@pytest.fixture(autouse=True)
def frozen_time():
    with time_machine.travel(NOW, tick=False):
        yield


@pytest.fixture
def conn(tmp_path):
    with db.connection(tmp_path / "test.db") as conn:
        clients.create_client(conn, "TechForce", pay_rate=35, aliases=["TFA"])
        clients.create_client(conn, "Patea", pay_rate=60)
        clients.create_project(conn, 1, "Website", pay_rate=50)
        clients.create_project(conn, 2, "Website")
        yield conn


def add(conn, start, end, **kwargs):
    entry_id, _ = clock.add_entry(
        conn, clock.parse_time(start), clock.parse_time(end), **kwargs
    )
    return entry_id


def ids(rep):
    return [row["id"] for row in rep.rows]


# --- Range and filters ---


def test_default_range_is_current_week(conn):
    add(conn, "2026-09-27 10:00", "2026-09-27 11:00")  # last Sunday
    add(conn, "2026-09-28 10:00", "2026-09-28 11:00")  # this Monday
    rep = report.build_report(conn)
    assert (rep.start, rep.end) == (date(2026, 9, 28), date(2026, 10, 4))
    assert ids(rep) == [2]


def test_end_date_is_inclusive_and_start_only_runs_to_today(conn):
    add(conn, "2026-09-10 23:30", "2026-09-10 23:45")
    add(conn, "2026-09-11 00:00", "2026-09-11 01:00")
    rep = report.build_report(conn, start=date(2026, 9, 1), end=date(2026, 9, 10))
    assert ids(rep) == [1]
    assert report.build_report(conn, start=date(2026, 9, 1)).end == date(2026, 9, 29)


def test_entry_belongs_to_range_it_starts_in(conn):
    add(conn, "2026-09-27 23:00", "2026-09-28 01:00")
    rep = report.build_report(conn, start=date(2026, 9, 21), end=date(2026, 9, 27))
    assert ids(rep) == [1]
    assert rep.total_seconds == 2 * 3600
    assert report.build_report(conn, start=date(2026, 9, 28)).rows == []


def test_open_entries_are_skipped_and_counted(conn):
    add(conn, "2026-09-29 08:00", "2026-09-29 09:00")
    clock.clock_in(conn, clock.parse_time("10:00"))
    rep = report.build_report(conn, start=date(2026, 9, 29))
    assert ids(rep) == [1]
    assert rep.skipped_open == 1


def test_client_filter_by_alias(conn):
    add(conn, "09:00", "10:00", client="TFA")
    add(conn, "10:00", "11:00", client="Patea")
    rep = report.build_report(conn, clients=["tfa"], start=date(2026, 9, 29))
    assert ids(rep) == [1]


def test_unknown_names_match_nothing(conn):
    add(conn, "09:00", "10:00", client="TFA")
    rep = report.build_report(conn, clients=["Nobody"], start=date(2026, 9, 29))
    assert rep.rows == []
    rep = report.build_report(conn, projects=["Nothing"], start=date(2026, 9, 29))
    assert rep.rows == []


def test_shared_project_name_matches_all_unless_client_given(conn):
    add(conn, "09:00", "10:00", client="TFA", project="Website")
    add(conn, "10:00", "11:00", client="Patea", project="Website")
    both = report.build_report(conn, projects=["website"], start=date(2026, 9, 29))
    assert ids(both) == [1, 2]
    one = report.build_report(
        conn, clients=["Patea"], projects=["website"], start=date(2026, 9, 29)
    )
    assert ids(one) == [2]


# --- Grouping ---


def test_group_by_day_splits_at_midnight(conn):
    add(conn, "2026-09-27 22:00", "2026-09-28 01:17:53", client="TFA")
    rep = report.build_report(
        conn, start=date(2026, 9, 27), end=date(2026, 9, 27), group_by="day"
    )
    assert [(r["period"], r["duration"], r["pay"]) for r in rep.rows] == [
        ("2026-09-27", 7200, 7000),
        ("2026-09-28", 4673, 4543),
    ]
    # Pieces add back up to the entry's stored total exactly
    assert rep.total_cents == 11543


def test_group_by_week_and_month_labels(conn):
    add(conn, "2026-09-23 09:00", "2026-09-23 10:00")
    week = report.build_report(conn, start=date(2026, 9, 1), group_by="week")
    assert week.rows[0]["period"] == "2026-09-21 to 2026-09-27"
    month = report.build_report(conn, start=date(2026, 9, 1), group_by="month")
    assert month.rows[0]["period"] == "2026-09"


def test_grouped_row_collects_lists_and_descriptions(conn):
    add(conn, "08:00", "09:00", client="TFA", description="First")
    add(conn, "10:00", "11:00", client="Patea", description="Second")
    add(conn, "11:00", "12:00", client="TFA")
    rep = report.build_report(conn, start=date(2026, 9, 29), group_by="day")
    row = rep.rows[0]
    assert row["id"] == [1, 2, 3]
    assert row["client"] == ["TechForce", "Patea"]
    assert row["description"] == "First\n\nSecond"
    assert row["duration"] == 3 * 3600


def test_days_without_work_have_no_row(conn):
    add(conn, "2026-09-21 09:00", "2026-09-21 10:00")
    add(conn, "2026-09-23 09:00", "2026-09-23 10:00")
    rep = report.build_report(conn, start=date(2026, 9, 21), group_by="day")
    assert [r["period"] for r in rep.rows] == ["2026-09-21", "2026-09-23"]


def test_no_rate_entries_have_no_pay(conn):
    add(conn, "09:00", "10:00")
    rep = report.build_report(conn, start=date(2026, 9, 29), group_by="day")
    assert rep.rows[0]["pay"] is None
    assert rep.total_cents is None


def test_split_uses_real_hours_on_dst_day():
    with pendulum.test_local_timezone(pendulum.timezone("America/New_York")):
        start = pendulum.datetime(2026, 11, 1, 0, 0, tz="America/New_York")
        end = pendulum.datetime(2026, 11, 2, 1, 0, tz="America/New_York")
        pieces = report.split_periods(start.int_timestamp, end.int_timestamp, "day")
    # Nov 1 has 25 hours when clocks fall back
    assert pieces == [(date(2026, 11, 1), 25 * 3600), (date(2026, 11, 2), 3600)]


@pytest.mark.parametrize(
    "total, seconds",
    [(100, [1, 1, 1]), (11543, [7200, 4673]), (1, [5, 5]), (0, [10, 20]), (99, [0])],
)
def test_allocate_always_sums_to_total(total, seconds):
    shares = report.allocate_cents(total, seconds)
    assert sum(shares) == total
    assert len(shares) == len(seconds)


# --- Fields ---


def test_fields_select_columns(conn):
    rep = report.build_report(conn, fields=["project", "pay"])
    assert rep.fields == ["project", "pay"]


def test_grouped_fields_always_start_with_period(conn):
    rep = report.build_report(conn, group_by="week", fields=["pay"])
    assert rep.fields == ["period", "pay"]


def test_entry_only_field_when_grouped_errors(conn):
    with pytest.raises(TimeTrackerError, match="only with --group-by entry"):
        report.build_report(conn, group_by="day", fields=["start"])


def test_unknown_field_errors(conn):
    with pytest.raises(TimeTrackerError, match="Unknown field"):
        report.build_report(conn, fields=["nope"])


# --- Formats ---


@pytest.fixture
def described(conn):
    add(conn, "08:00", "09:00", client="TFA", description="A | B\n\nC\tD")
    return report.build_report(
        conn, start=date(2026, 9, 29), fields=["id", "description", "duration", "pay"]
    )


def test_markdown_escapes_pipes_and_newlines(described):
    lines = render_markdown(described).splitlines()
    assert lines[2] == "| 1 | A \\| B<br><br>C\tD | 1:00:00 | $35.00 |"
    assert lines[3] == "| Total |  | 1:00:00 | $35.00 |"


def test_csv_quotes_multiline_and_keeps_numbers_plain(described):
    rows = list(csv.reader(io.StringIO(render_csv(described))))
    assert rows[1] == ["1", "A | B\n\nC\tD", "1:00:00", "35.00"]


def test_tsv_escapes_so_each_line_is_a_row(described):
    lines = render_tsv(described).splitlines()
    assert len(lines) == 2
    assert lines[1] == "1\tA | B\\n\\nC\\tD\t1:00:00\t35.00"


def test_json_uses_raw_values(conn):
    add(conn, "08:00", "09:30", client="TFA")
    rep = report.build_report(conn, start=date(2026, 9, 29))
    data = json.loads(render_json(rep))
    row = data["rows"][0]
    assert row["start"] == clock.parse_time("08:00").isoformat()
    assert row["duration"] == 5400
    assert row["pay"] == 52.5
    assert data["totals"] == {"duration": 5400, "pay": 52.5}


# --- CLI ---


@pytest.fixture
def cli(tmp_path, monkeypatch):
    monkeypatch.setenv("TIMETRACKER_DB", str(tmp_path / "data" / "cli.db"))
    monkeypatch.chdir(tmp_path)
    runner = CliRunner()
    invoke = lambda *args: runner.invoke(app, list(args))  # noqa: E731
    invoke("client", "add", "Acme", "--pay-rate-hourly", "40")
    invoke("add", "--start-time", "08:00", "--end-time", "09:00", "--client", "Acme")
    return invoke


def test_cli_prints_csv_to_stdout(cli):
    result = cli("report", "--output", "csv", "--fields", "id,pay")
    assert result.stdout == "ID,Pay\n1,40.00\n"


def test_cli_write_uses_reports_folder(cli, tmp_path):
    result = cli("report", "--output", "json", "--write")
    assert result.exit_code == 0
    saved = list((tmp_path / "data" / "reports").iterdir())
    assert [p.name for p in saved] == ["report-2026-09-28-to-2026-10-04-by-entry.json"]


def test_cli_filename_with_folder_is_used_as_given(cli, tmp_path):
    result = cli("report", "--output", "md", "--filename", "out/hours.md")
    assert result.exit_code == 0
    assert (tmp_path / "out" / "hours.md").read_text().startswith("| ID |")


def test_cli_table_cannot_be_written(cli):
    result = cli("report", "--write")
    assert result.exit_code == 1
    assert "--output md" in result.output


def test_cli_warns_about_open_entries(cli):
    cli("in", "--client", "Acme")
    result = cli("report")
    assert "1 open entry not included" in result.output


@pytest.mark.parametrize(
    "args", [("--group-by", "year"), ("--output", "docx"), ("--start", "9/1")]
)
def test_cli_rejects_bad_options(cli, args):
    result = cli("report", *args)
    assert result.exit_code == 1


# --- PDF and Excel ---


def test_report_xlsx_keeps_values_typed(conn, tmp_path):
    from timetracker.formats.xlsx import render_report_xlsx

    add(conn, "08:00", "09:30", client="TFA", description="A\n\nB")
    rep = report.build_report(conn, start=date(2026, 9, 29))
    path = tmp_path / "report.xlsx"
    render_report_xlsx(rep, path)

    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None
    assert sheet["A1"].value == "TIME REPORT"
    assert sheet["A2"].value == "2026-09-29 to 2026-09-29, by entry"

    headers = [c.value for c in sheet[4]]
    assert headers == [
        "ID",
        "Start",
        "End",
        "Client",
        "Project",
        "Description",
        "Duration",
        "Rate",
        "Pay",
    ]

    duration = sheet.cell(row=5, column=headers.index("Duration") + 1)
    assert duration.value == timedelta(seconds=5400)
    assert duration.number_format == "[h]:mm:ss"
    pay = sheet.cell(row=5, column=headers.index("Pay") + 1)
    assert pay.value == 52.5
    start = sheet.cell(row=5, column=headers.index("Start") + 1)
    assert start.value == datetime(2026, 9, 29, 8, 0)


def test_report_xlsx_total_lines_up_with_its_columns(conn, tmp_path):
    from timetracker.formats.xlsx import render_report_xlsx

    add(conn, "08:00", "09:00", client="TFA")
    add(conn, "10:00", "11:00", client="TFA")
    rep = report.build_report(
        conn,
        start=date(2026, 9, 29),
        group_by="day",
        fields=["period", "duration", "pay"],
    )
    path = tmp_path / "grouped.xlsx"
    render_report_xlsx(rep, path)

    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None
    total = sheet[sheet.max_row]
    assert [c.value for c in total] == [
        "Total",
        timedelta(seconds=7200),
        70.0,
    ]


def test_report_pdf_is_written_and_turns_landscape(conn, tmp_path):
    from timetracker.formats.pdf import _report_source, render_report_pdf

    add(conn, "08:00", "09:00", client="TFA", description="Piped | text")
    wide = report.build_report(conn, start=date(2026, 9, 29))
    narrow = report.build_report(
        conn, start=date(2026, 9, 29), fields=["id", "duration"]
    )

    assert "flipped: true" in _report_source(wide)
    assert "flipped" not in _report_source(narrow)

    path = tmp_path / "report.pdf"
    render_report_pdf(wide, path)
    assert path.read_bytes().startswith(b"%PDF")


def test_report_pdf_keeps_multiline_descriptions(conn, tmp_path):
    from timetracker.formats.pdf import _report_source

    add(conn, "08:00", "09:00", client="TFA", description="First\n\nSecond")
    rep = report.build_report(conn, start=date(2026, 9, 29), fields=["description"])
    source = _report_source(rep)
    assert "First \\\nSecond" in source


def test_cli_pdf_and_xlsx_write_without_the_write_flag(cli, tmp_path):
    assert cli("report", "--output", "xlsx").exit_code == 0
    assert (
        cli("report", "--output", "pdf", "--filename", "out/hours.pdf").exit_code == 0
    )

    saved = list((tmp_path / "data" / "reports").iterdir())
    assert [p.name for p in saved] == ["report-2026-09-28-to-2026-10-04-by-entry.xlsx"]
    assert (tmp_path / "out" / "hours.pdf").exists()
